
import io
import json
import random
import string
import time
from datetime import datetime, timedelta

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from flask import flash, jsonify, redirect, render_template, request, send_file, session, url_for
from payos.type import ItemData, PaymentData

from decorators import passenger_login_required
from extensions import db, payos_client
from models import Booking, FoodOrder, MenuItem, RestStop, RestStopReview, Trip, User


INVALID_TRIP_STATUSES = {'Cancelled', 'Completed', 'Broken'}


def generate_ticket_code():
    while True:
        prefix = ''.join(random.choices(string.ascii_uppercase, k=2))
        numbers = ''.join(random.choices(string.digits, k=8))
        code = prefix + numbers
        if not Booking.query.filter_by(ticket_code=code).first():
            return code


def _now():
    return datetime.now()


def _to_base36(num):
    digits = '0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ'
    num = int(num)
    if num == 0:
        return '0'
    out = []
    while num:
        num, rem = divmod(num, 36)
        out.append(digits[rem])
    return ''.join(reversed(out))


def _build_archived_seat_number(original_seat, booking_id):
    original = (original_seat or '').strip() or 'SEAT'
    encoded = _to_base36(booking_id)
    archived = f'~{original}~{encoded}'
    if len(archived) <= 10:
        return archived
    original_short = original[:3] if original else 'S'
    encoded_short = encoded[-5:]
    return f'~{original_short}~{encoded_short}'[:10]


def _extract_display_seat_number(seat_number):
    seat_str = str(seat_number or '')
    if seat_str.startswith('~') and seat_str.count('~') >= 2:
        parts = seat_str.split('~')
        if len(parts) >= 3 and parts[1]:
            return parts[1]
    return seat_str


def _normalize_booking_for_display(booking):
    if booking is not None:
        booking.seat_number = _extract_display_seat_number(booking.seat_number)
    return booking


def _trip_accepts_new_bookings(trip, now=None):
    now = now or _now()
    if trip is None:
        return False
    if trip.status in INVALID_TRIP_STATUSES:
        return False
    if trip.departure_time <= now:
        return False
    return True


def _is_active_hold(booking, now=None):
    now = now or _now()
    return (
        booking.status == 'HOLD'
        and booking.expires_at is not None
        and booking.expires_at > now
    )


def _can_view_ticket(booking):
    return booking.status == 'CONFIRMED' and booking.payment_status == 'Paid'


def _archive_booking_for_reuse(booking, reason='EXPIRED'):
    if booking is None:
        return False

    if str(booking.seat_number).startswith('~') and booking.status == 'CANCELLED':
        booking.expires_at = None
        return False

    original_seat = _extract_display_seat_number(booking.seat_number)
    booking.seat_number = _build_archived_seat_number(original_seat, booking.id)
    booking.expires_at = None
    booking.status = 'CANCELLED'

    if reason == 'EXPIRED':
        if booking.payment_status != 'Paid':
            booking.payment_status = 'Expired'
        booking.payment_id = None
    elif reason == 'CANCELLED':
        if booking.payment_status == 'Paid':
            booking.payment_status = 'Refund Pending'
        else:
            booking.payment_status = 'Cancelled'
        booking.payment_id = None

    return True


def _cleanup_expired_holds_for_trip(trip_id):
    now = _now()
    stale_holds = Booking.query.filter(
        Booking.trip_id == trip_id,
        Booking.status == 'HOLD',
        Booking.expires_at.isnot(None),
        Booking.expires_at <= now
    ).all()

    changed = False
    for booking in stale_holds:
        changed = _archive_booking_for_reuse(booking, reason='EXPIRED') or changed

    if changed:
        db.session.commit()

    return changed


def _cleanup_expired_holds_for_user(user_id):
    now = _now()
    stale_holds = Booking.query.filter(
        Booking.user_id == user_id,
        Booking.status == 'HOLD',
        Booking.expires_at.isnot(None),
        Booking.expires_at <= now
    ).all()

    changed = False
    for booking in stale_holds:
        changed = _archive_booking_for_reuse(booking, reason='EXPIRED') or changed

    if changed:
        db.session.commit()

    return changed


def _release_reusable_seat_rows(trip_id, seat_number):
    now = _now()
    reusable_rows = Booking.query.filter(
        Booking.trip_id == trip_id,
        Booking.seat_number == seat_number
    ).all()

    changed = False
    for booking in reusable_rows:
        is_expired_hold = booking.status == 'HOLD' and booking.expires_at is not None and booking.expires_at <= now
        is_cancelled = booking.status == 'CANCELLED'
        if is_expired_hold:
            changed = _archive_booking_for_reuse(booking, reason='EXPIRED') or changed
        elif is_cancelled:
            changed = _archive_booking_for_reuse(booking, reason='CANCELLED') or changed

    if changed:
        db.session.commit()

    return changed


def _ensure_booking_is_payable(booking, archive_expired=True):
    now = _now()

    if booking.status == 'CANCELLED':
        return False, 'Vé này đã bị hủy, không thể thanh toán.', 'my_tickets'

    if booking.payment_status == 'Paid':
        return False, 'Vé này đã được thanh toán trước đó.', 'view_ticket'

    if booking.status != 'HOLD':
        return False, 'Vé này không còn ở trạng thái chờ thanh toán.', 'my_tickets'

    if booking.expires_at is None or booking.expires_at <= now:
        if archive_expired:
            changed = _archive_booking_for_reuse(booking, reason='EXPIRED')
            if changed:
                db.session.commit()
        return False, 'Vé giữ chỗ đã hết hạn. Vui lòng đặt lại.', 'my_tickets'

    if not _trip_accepts_new_bookings(booking.trip, now):
        return False, 'Chuyến đi này không còn hợp lệ để thanh toán.', 'my_tickets'

    return True, None, None


def _confirm_paid_booking(booking, method):
    booking.status = 'CONFIRMED'
    booking.payment_status = 'Paid'
    booking.payment_method = method
    booking.expires_at = None


def register_booking_routes(app):
    def book_seat(trip_id):
        if 'user_id' not in session and 'company_id' not in session:
            flash('Vui lòng đăng nhập!')
            return redirect(url_for('passenger_login'))

        trip = Trip.query.get_or_404(trip_id)
        now = _now()

        if not _trip_accepts_new_bookings(trip, now):
            flash('Chuyến đi này không còn nhận đặt vé.')
            return redirect(url_for('index'))

        _cleanup_expired_holds_for_trip(trip_id)

        bookings = Booking.query.filter_by(trip_id=trip_id, status='CONFIRMED').all()
        occupied_seats = [b.seat_number for b in bookings]

        holds = Booking.query.filter(
            Booking.trip_id == trip_id,
            Booking.status == 'HOLD',
            Booking.expires_at > now
        ).all()
        held_seats = [b.seat_number for b in holds]

        user = db.session.get(User, session.get('user_id')) if 'user_id' in session else None
        seat_map = json.loads(trip.bus.seat_map)
        stops = RestStop.query.all()

        return render_template(
            'book.html',
            trip=trip,
            seat_map=seat_map,
            occupied_seats=occupied_seats,
            held_seats=held_seats,
            user=user,
            stops=stops
        )

    def process_booking():
        if 'user_id' not in session and 'company_id' not in session:
            return 'Unauthorized', 401

        trip_id = request.form.get('trip_id')
        seat_number = (request.form.get('seat_number') or '').strip()
        passenger_name = (request.form.get('name') or '').strip()
        passenger_phone = (request.form.get('phone') or '').strip()
        pickup_point = request.form.get('pickup_point')
        dropoff_point = request.form.get('dropoff_point')
        user_id = session.get('user_id')

        if not trip_id or not seat_number or not passenger_name or not passenger_phone:
            flash('Thiếu thông tin đặt vé.')
            return redirect(url_for('index'))

        trip = Trip.query.get_or_404(trip_id)
        now = _now()

        if not _trip_accepts_new_bookings(trip, now):
            flash('Chuyến đi này không còn nhận đặt vé.')
            return redirect(url_for('index'))

        _cleanup_expired_holds_for_trip(trip.id)
        _release_reusable_seat_rows(trip.id, seat_number)

        existing = Booking.query.filter(
            Booking.trip_id == trip.id,
            Booking.seat_number == seat_number,
            (
                (Booking.status == 'CONFIRMED')
                | ((Booking.status == 'HOLD') & (Booking.expires_at > now))
            )
        ).first()

        if existing:
            flash('Ghế này đã có người giữ hoặc đặt!')
            return redirect(url_for('book_seat', trip_id=trip.id))

        final_price = trip.route.base_price * trip.price_multiplier

        status = 'CONFIRMED' if 'company_id' in session else 'HOLD'
        expires_at = (now + timedelta(minutes=15)) if status == 'HOLD' else None
        payment_status = 'Paid' if status == 'CONFIRMED' else 'Unpaid'
        payment_method = 'Tiền mặt' if status == 'CONFIRMED' else None

        ticket_code = generate_ticket_code()

        new_booking = Booking(
            trip_id=trip.id,
            seat_number=seat_number,
            user_id=user_id,
            passenger_name=passenger_name,
            passenger_phone=passenger_phone,
            pickup_point=pickup_point,
            dropoff_point=dropoff_point,
            status=status,
            expires_at=expires_at,
            ticket_price=final_price,
            payment_status=payment_status,
            payment_method=payment_method,
            ticket_code=ticket_code
        )

        try:
            db.session.add(new_booking)
            db.session.flush()

            food_total = 0
            food_details = {}
            for key, val in request.form.items():
                if key.startswith('qty_') and val.isdigit() and int(val) > 0:
                    item_id = int(key.split('_')[1])
                    qty = int(val)
                    item = MenuItem.query.get(item_id)
                    if item:
                        cost = item.price * qty
                        food_total += cost
                        food_details[item.name] = {'qty': qty, 'price': item.price}

            if food_total > 0:
                food_order = FoodOrder(
                    booking_id=new_booking.id,
                    order_details=json.dumps(food_details),
                    total_price=food_total
                )
                db.session.add(food_order)

            db.session.commit()
        except Exception:
            db.session.rollback()
            flash('Lỗi đặt vé: đã có người đặt nhanh hơn bạn!')
            return redirect(url_for('book_seat', trip_id=trip.id))

        if 'company_id' in session:
            flash(f'Đã đặt vé thành công! Mã vé: #{new_booking.id}')
            return redirect(url_for('admin_trips_page'))

        flash('Đã giữ vé thành công! Vui lòng thanh toán trong 15 phút.')
        return redirect(url_for('payment', booking_id=new_booking.id))

    def payment(booking_id):
        booking = Booking.query.get_or_404(booking_id)

        if booking.user_id != session.get('user_id'):
            flash('Không tìm thấy đơn đặt vé!')
            return redirect(url_for('index'))

        allowed, message, endpoint = _ensure_booking_is_payable(booking, archive_expired=True)
        if not allowed:
            flash(message)
            if endpoint == 'view_ticket' and _can_view_ticket(booking):
                return redirect(url_for('view_ticket', booking_id=booking.id))
            return redirect(url_for(endpoint))

        _normalize_booking_for_display(booking)

        if request.method == 'POST':
            method = (request.form.get('payment_method') or '').strip()
            if not method:
                flash('Vui lòng chọn phương thức thanh toán.')
                return redirect(url_for('payment', booking_id=booking_id))

            allowed, message, endpoint = _ensure_booking_is_payable(booking, archive_expired=True)
            if not allowed:
                flash(message)
                return redirect(url_for(endpoint))

            if method == 'PayOS':
                if payos_client is None:
                    flash('Cổng thanh toán PayOS chưa được cấu hình.')
                    return redirect(url_for('payment', booking_id=booking_id))

                domain = request.host_url.rstrip('/')
                item = ItemData(
                    name=f"Ve xe {booking.trip.route.start_point} - {booking.trip.route.end_point}",
                    quantity=1,
                    price=int(booking.ticket_price)
                )

                order_code = int(time.time() * 1000) + int(booking.id)
                booking.payment_id = str(order_code)
                booking.payment_method = 'PayOS'
                db.session.commit()

                payment_data = PaymentData(
                    orderCode=order_code,
                    amount=int(booking.ticket_price),
                    description=f"Thanh toan ve {booking.ticket_code}",
                    items=[item],
                    cancelUrl=f"{domain}/payment/cancel/{booking_id}",
                    returnUrl=f"{domain}/payment/success/{booking_id}"
                )

                try:
                    response = payos_client.createPaymentLink(payment_data)
                    return redirect(response.checkoutUrl)
                except Exception as e:
                    print(f'PayOS Error: {e}')
                    flash('Có lỗi xảy ra khi kết nối với cổng thanh toán PayOS. Vui lòng thử lại sau.')
                    return redirect(url_for('payment', booking_id=booking_id))

            _confirm_paid_booking(booking, method)
            db.session.commit()
            flash('Thanh toán thành công! Vé của bạn đã được xuất.')
            return redirect(url_for('view_ticket', booking_id=booking.id))

        return render_template('payment.html', booking=booking)

    def payment_success(booking_id):
        booking = Booking.query.get_or_404(booking_id)

        if booking.user_id != session.get('user_id'):
            flash('Không tìm thấy đơn đặt vé!')
            return redirect(url_for('index'))

        if _can_view_ticket(booking):
            flash('Vé này đã được thanh toán thành công trước đó.')
            return redirect(url_for('view_ticket', booking_id=booking.id))

        allowed, message, endpoint = _ensure_booking_is_payable(booking, archive_expired=True)
        if not allowed:
            flash(message)
            return redirect(url_for(endpoint))

        try:
            if payos_client is None:
                flash('Cổng thanh toán PayOS chưa được cấu hình.')
                return redirect(url_for('payment', booking_id=booking.id))

            if not booking.payment_id:
                flash('Không tìm thấy thông tin thanh toán.')
                return redirect(url_for('payment', booking_id=booking.id))

            payment_info = payos_client.getPaymentLinkInformation(int(booking.payment_id))
            if getattr(payment_info, 'status', None) == 'PAID':
                _confirm_paid_booking(booking, 'PayOS')
                db.session.commit()
                flash('Thanh toán thành công qua PayOS! Vé của bạn đã được xuất.')
                return redirect(url_for('view_ticket', booking_id=booking.id))

            flash('Thanh toán chưa hoàn tất hoặc thất bại.')
            return redirect(url_for('payment', booking_id=booking.id))
        except Exception as e:
            print(f'PayOS Check Error: {e}')
            flash('Không thể xác minh trạng thái thanh toán lúc này.')
            return redirect(url_for('payment', booking_id=booking.id))

    def payment_cancel(booking_id):
        booking = Booking.query.get_or_404(booking_id)

        if booking.user_id != session.get('user_id'):
            flash('Không tìm thấy đơn đặt vé!')
            return redirect(url_for('index'))

        allowed, message, endpoint = _ensure_booking_is_payable(booking, archive_expired=True)
        if not allowed:
            flash(message)
            return redirect(url_for(endpoint))

        flash('Bạn đã hủy thanh toán qua PayOS.')
        return redirect(url_for('payment', booking_id=booking_id))

    def payos_webhook():
        data = request.get_json(silent=True)
        if not data:
            return jsonify({'error': 'No data'}), 400

        try:
            if payos_client is None:
                return jsonify({'error': 'PayOS is not configured'}), 503

            try:
                verified = payos_client.webhooks.verify(data)
            except Exception:
                verified = payos_client.verifyPaymentWebhookData(data)

            order_code = str(getattr(verified, 'orderCode', ''))
            amount = int(getattr(verified, 'amount', 0) or 0)
            result_code = str(getattr(verified, 'code', ''))

            if not order_code:
                return jsonify({'error': 'Missing orderCode'}), 400

            bookings = Booking.query.filter_by(payment_id=order_code).all()
            if not bookings:
                return jsonify({'error': 'Booking not found'}), 404

            total_ticket_price = sum(int(b.ticket_price) for b in bookings)
            if amount and total_ticket_price != amount:
                return jsonify({'error': 'Amount mismatch'}), 400

            if result_code not in ('00', '0'):
                return jsonify({'success': True, 'message': 'Payment not successful'}), 200

            for booking in bookings:
                allowed, _, _ = _ensure_booking_is_payable(booking, archive_expired=True)
                if not allowed:
                    if _can_view_ticket(booking):
                        continue
                    else:
                        continue # Skip unpayable ones

                if booking.payment_status != 'Paid':
                    _confirm_paid_booking(booking, 'PayOS')
                    print(f'Webhook Success: Booking {booking.id} confirmed.')

            db.session.commit()

            return jsonify({'success': True}), 200
        except Exception as e:
            print(f'Webhook Error: {e}')
            return jsonify({'error': str(e)}), 400

    def view_meal_ticket(booking_id):
        booking = Booking.query.get_or_404(booking_id)
        if booking.user_id != session['user_id']:
            flash('Không có quyền truy cập!')
            return redirect(url_for('index'))

        food_order = FoodOrder.query.filter_by(booking_id=booking_id).first()
        if not food_order:
            flash('Bạn chưa đặt suất ăn cho vé này.')
            return redirect(url_for('my_tickets'))

        _normalize_booking_for_display(booking)
        return render_template('meal_ticket.html', booking=booking, order=food_order)

    def download_ticket_word(booking_id):
        booking = Booking.query.get_or_404(booking_id)
        if booking.user_id != session['user_id']:
            flash('Unauthorized')
            return redirect(url_for('index'))

        if not _can_view_ticket(booking):
            flash('Chỉ có thể tải vé khi đã thanh toán thành công.')
            return redirect(url_for('my_tickets'))

        _normalize_booking_for_display(booking)

        document = Document()
        document.add_heading('VÉ XE ĐIỆN TỬ - HUTECH BUS', 0)

        p = document.add_paragraph()
        p.add_run(f'Mã vé: #{booking.id}').bold = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        table = document.add_table(rows=6, cols=2)
        table.style = 'Table Grid'

        def add_row(idx, label, value):
            row = table.rows[idx]
            row.cells[0].text = label
            row.cells[1].text = str(value)

        add_row(0, 'Hành khách', booking.passenger_name)
        add_row(1, 'SĐT', booking.passenger_phone)
        add_row(2, 'Tuyến', f'{booking.trip.route.start_point} - {booking.trip.route.end_point}')
        add_row(3, 'Khởi hành', booking.trip.departure_time.strftime('%H:%M %d/%m/%Y'))
        add_row(4, 'Biển số xe', booking.trip.bus.license_plate)
        add_row(5, 'Số ghế', booking.seat_number)

        document.add_paragraph().add_run(f'\nGiá vé: {"{:,.0f}".format(booking.ticket_price)} VNĐ').bold = True
        document.add_paragraph(f'Trạng thái: {booking.payment_status}').italic = True
        document.add_paragraph('\nCảm ơn quý khách đã sử dụng dịch vụ Hutech Bus!')

        f = io.BytesIO()
        document.save(f)
        f.seek(0)

        return send_file(
            f,
            as_attachment=True,
            download_name=f'ticket_{booking.id}.docx',
            mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        )

    def food_order(booking_id):
        booking = Booking.query.get_or_404(booking_id)
        if booking.user_id != session['user_id']:
            flash('Unauthorized')
            return redirect(url_for('index'))

        stops = RestStop.query.all()

        if request.method == 'POST':
            total = 0
            details = {}

            for key, val in request.form.items():
                if key.startswith('qty_') and val.isdigit() and int(val) > 0:
                    item_id = int(key.split('_')[1])
                    qty = int(val)
                    item = MenuItem.query.get(item_id)
                    if item:
                        cost = item.price * qty
                        total += cost
                        details[item.name] = {'qty': qty, 'price': item.price}

            if total > 0:
                order = FoodOrder(booking_id=booking_id, order_details=json.dumps(details), total_price=total)
                db.session.add(order)
                db.session.commit()
                flash('Đặt món thành công! Vui lòng thanh toán khi nhận món.')
                return redirect(url_for('my_tickets'))

            flash('Vui lòng chọn ít nhất 1 món!')

        existing_order = FoodOrder.query.filter_by(booking_id=booking_id).first()
        _normalize_booking_for_display(booking)
        return render_template('food_order.html', booking=booking, stops=stops, existing_order=existing_order)

    def rate_stop():
        booking_id = request.form.get('booking_id')
        stop_id = request.form.get('stop_id')
        rating = int(request.form.get('rating'))
        comment = request.form.get('comment')

        new_review = RestStopReview(booking_id=booking_id, rest_stop_id=stop_id, rating=rating, comment=comment)
        db.session.add(new_review)
        db.session.commit()
        flash('Cảm ơn bạn đã đánh giá trạm dừng!')
        return redirect(url_for('my_tickets'))

    app.add_url_rule('/book/<int:trip_id>', view_func=book_seat)
    app.add_url_rule('/booking/confirm', view_func=process_booking, methods=['POST'])
    app.add_url_rule('/payment/<int:booking_id>', view_func=passenger_login_required(payment), methods=['GET', 'POST'])
    app.add_url_rule('/payment/success/<int:booking_id>', view_func=passenger_login_required(payment_success))
    app.add_url_rule('/payment/cancel/<int:booking_id>', view_func=passenger_login_required(payment_cancel))
    app.add_url_rule('/payos-webhook', view_func=payos_webhook, methods=['POST'])
    app.add_url_rule('/ticket/meal/<int:booking_id>', view_func=passenger_login_required(view_meal_ticket))
    app.add_url_rule('/ticket/<int:booking_id>/word', view_func=passenger_login_required(download_ticket_word))
    app.add_url_rule('/booking/<int:booking_id>/food', view_func=passenger_login_required(food_order), methods=['GET', 'POST'])
    app.add_url_rule('/api/rate_stop', view_func=passenger_login_required(rate_stop), methods=['POST'])
